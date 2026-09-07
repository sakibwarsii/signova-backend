import io
import asyncio
import fitz  # PyMuPDF
import docx
from pptx import Presentation
from fastapi import UploadFile

def extract_text_from_pdf(file_bytes: bytes, max_pages: int = 10) -> str:
    text = ""
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    pages_to_process = min(len(doc), max_pages)
    for i in range(pages_to_process):
        page = doc.load_page(i)
        text += page.get_text() + "\n"
    return text

def extract_text_from_docx(file_bytes: bytes, max_paras: int = 200) -> str:
    text = ""
    doc = docx.Document(io.BytesIO(file_bytes))
    paras_to_process = min(len(doc.paragraphs), max_paras)
    for i in range(paras_to_process):
        text += doc.paragraphs[i].text + "\n"
    return text

def extract_text_from_pptx(file_bytes: bytes, max_slides: int = 20) -> str:
    text = ""
    prs = Presentation(io.BytesIO(file_bytes))
    slides_to_process = min(len(prs.slides), max_slides)
    for i in range(slides_to_process):
        slide = prs.slides[i]
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"
    return text

async def extract_text(file: UploadFile, max_words: int = 5000) -> str:
    """
    Extracts text from the uploaded file and truncates it to max_words
    to prevent LLM token exhaustion.
    """
    file_bytes = await file.read()
    filename = file.filename.lower()

    # Parsing (PyMuPDF/python-docx/python-pptx) is CPU/disk-bound and synchronous —
    # run it off the event loop so it doesn't stall other requests/websockets
    # while a large PDF/PPTX is being parsed.
    text = ""
    if filename.endswith(".pdf"):
        text = await asyncio.to_thread(extract_text_from_pdf, file_bytes)
    elif filename.endswith(".docx"):
        text = await asyncio.to_thread(extract_text_from_docx, file_bytes)
    elif filename.endswith(".pptx"):
        text = await asyncio.to_thread(extract_text_from_pptx, file_bytes)
    elif filename.endswith(".txt"):
        text = file_bytes.decode('utf-8', errors='ignore')
    else:
        raise ValueError("Unsupported file format")

    # Limit to max_words to save tokens
    words = text.split()
    if len(words) > max_words:
        words = words[:max_words]
        text = " ".join(words)
    
    return text

def chunk_text(text: str, chunk_size_words: int = 200) -> list[str]:
    """
    Splits a large text into smaller chunks of approximately chunk_size_words.
    """
    words = text.split()
    chunks = []
    for i in range(0, len(words), chunk_size_words):
        chunks.append(" ".join(words[i:i + chunk_size_words]))
    return chunks
